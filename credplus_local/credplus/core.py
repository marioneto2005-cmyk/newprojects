"""Local extraction, explicit output allowlist and reviewed document generation."""
from __future__ import annotations
import html, json, os, re, shutil, subprocess, sys, tempfile
from dataclasses import dataclass
from pathlib import Path
import fitz
from .rule_engine import extract, valid_id

ASSETS = Path(__file__).resolve().parent / 'assets'
INSURERS = {'Zurich': ['Única'], 'HDI': ['Opção 1 (confira a modalidade no PDF)', 'Opção 2 (confira a modalidade no PDF)'],
            'Allianz': ['Básico', 'Ampliado', 'Completo', 'Master', 'Exclusivo'],
            'Tokio Marine': ['Opção 1', 'Opção 2'], 'Yelum': ['Única']}
@dataclass(frozen=True)
class Field:
    key: str
    label: str
    section: str
    kind: str = 'text'
    cell: str = ''
FIELDS = [Field(*x) for x in [
 ('nome','Nome do proponente','Segurado','text','B5'),('cpf','CPF / CNPJ','Segurado','id','D5'),
 ('condutor','Principal condutor','Segurado','text','B6'),('cpfcond','CPF / CNPJ do condutor','Segurado','id','D6'),
 ('validade','Validade da proposta','Segurado','date','D7'),('inicio','Início da vigência','Segurado','date',''),('fim','Fim da vigência','Segurado','date',''),
 ('uso','Uso do veículo','Perfil','text','B10'),('trabalho','Uso para trabalho','Perfil','text','D10'),
 ('faculdade','Uso para faculdade','Perfil','text','B11'),('garagem','Garagem no pernoite','Perfil','text','D11'),
 ('jovens','Condutores jovens / faixa etária','Perfil','text','B12'),('cep','CEP de pernoite','Perfil','id','D12'),
 ('modelo','Marca / modelo','Veículo','text','B15'),('ano','Ano modelo / fabricação','Veículo','text','D15'),
 ('placa','Placa','Veículo','id','B16'),('chassi','Chassi','Veículo','id','D16'),
 ('combustivel','Combustível','Veículo','text','B17'),('zero','Zero km','Veículo','text','D17'),
 ('casco','Cobertura do casco / FIPE','Coberturas','text','C21'),('dm','Danos materiais a terceiros','Coberturas','money','C22'),
 ('dc','Danos corporais a terceiros','Coberturas','money','C23'),('appm','APP - morte','Coberturas','money','C24'),
 ('appi','APP - invalidez permanente','Coberturas','money','C25'),('appg','APP - garantia conjunta por passageiro','Coberturas','money',''),
 ('morais','Danos morais / estéticos','Coberturas','money','C26'),('vidros','Proteção a vidros','Coberturas','text','C27'),
 ('assistencia','Assistência 24 horas','Coberturas','text','C28'),('reserva','Carro reserva','Coberturas','text','C29'),
 ('franquia','Franquia do casco','Franquias','money','C31'),('tipofranquia','Tipo de franquia (conforme proposta)','Franquias','text',''),
 ('parabrisa','Para-brisa','Franquias','money','B32'),('laterais','Vidros laterais','Franquias','money','B33'),
 ('farol','Farol convencional','Franquias','money','B34'),('traseiro','Vidro traseiro','Franquias','money','B35'),
 ('farolled','Farol xenônio / LED','Franquias','money','B36'),('lanternaled','Lanterna LED','Franquias','money','D32'),
 ('retrovisor','Retrovisor','Franquias','money','D33'),('lanterna','Lanterna convencional','Franquias','money','D34'),
 ('auxiliar','Farol auxiliar / milha','Franquias','money','D35'),('lanternaaux','Lanterna auxiliar','Franquias','money','D36'),
 ('total','Prêmio total da opção selecionada','Pagamento','money','C38'),
 ('pagamento_cartao','Condições de pagamento - cartão','Pagamento','text','C39'),
 ('pagamento_outros','Condições de pagamento - débito / outros','Pagamento','text','C40'),
]]
# Optional reviewed complements cover commercial details outside the spreadsheet slots.
FIELDS += [Field(key,label,'Informações adicionais','supplement') for key,label in [
 ('segurado_extra','Outros dados do segurado / condutor'),
 ('perfil_extra','Outros dados do perfil de risco'),
 ('veiculo_extra','Outros dados do veículo'),
 ('coberturas_extra','Outras coberturas e limites'),
 ('franquias_extra','Outras franquias e condições'),
]]
FIELDS += [Field(*x) for x in [
 ('endereco','Endereço do segurado','Segurado','text',''),
 ('bairro','Bairro do segurado','Segurado','text',''),
 ('cidade_uf','Cidade / UF do segurado','Segurado','text',''),
 ('cep_segurado','CEP do segurado','Segurado','id',''),
 ('endereco_pernoite','Endereço de pernoite','Perfil','text',''),
 ('cidade_pernoite','Cidade / UF de pernoite','Perfil','text',''),
 ('cep_circulacao','CEP de circulação','Perfil','id',''),
]]
FIELD_MAP = {f.key:f for f in FIELDS}


def safe_text(value):
    """Keep names/brands as supplied; discard control characters only."""
    return ''.join(c for c in str(value).strip() if c in '\n\t' or ord(c)>=32)


def format_value(field, value):
    if value is None or str(value).strip()=='': return ''
    text = str(value).strip()
    if field.kind == 'money':
        if isinstance(value,(int,float)): n=float(value)
        else:
            t=text.replace('R$','').replace(' ','')
            if not re.fullmatch(r'(?:\d{1,3}(?:\.\d{3})*|\d+),\d{2}',t):
                raise ValueError(f'{field.label}: informe o valor em reais, como 1.234,56.')
            n=float(t.replace('.','').replace(',','.'))
        if not 0 <= n <= 1e12: raise ValueError(f'{field.label}: valor fora do intervalo permitido.')
        return f'{n:,.2f}'.replace(',','_').replace('.',',').replace('_','.')
    if field.key in ('cep','cep_segurado','cep_circulacao'):
        partial=re.fullmatch(r'(\d{5})(?:\s*\(parcial(?: no documento)?\))?',text,re.I)
        if partial:return partial[1]+' (parcial no documento)'
    if field.kind=='id' and not valid_id('cep' if field.key in ('cep_segurado','cep_circulacao') else field.key,text):
        raise ValueError(f'{field.label}: formato ou dígitos inválidos; confira o PDF. Não adivinhe caracteres.')
    if field.kind=='date':
        import datetime
        for fmt in ('%d/%m/%Y','%Y-%m-%d'):
            try:return datetime.datetime.strptime(text,fmt).strftime('%d/%m/%Y')
            except ValueError: pass
        raise ValueError(f'{field.label}: use DD/MM/AAAA.')
    if len(text)>1600: raise ValueError(f'{field.label}: texto longo demais; informe somente o dado correspondente.')
    return safe_text(text)


def validate_values(values):
    clean={f.key:format_value(f,values.get(f.key,'')) for f in FIELDS}
    for key in ('nome','modelo','total'):
        if not clean[key]: raise ValueError(f'Preencha {FIELD_MAP[key].label} antes de gerar o documento.')
    return clean


def _run(command, timeout=180):
    opts={'capture_output':True,'timeout':timeout}
    if os.name=='nt':opts['creationflags']=subprocess.CREATE_NO_WINDOW
    result=subprocess.run(command,**opts)
    if result.returncode:raise RuntimeError(f'A ferramenta local falhou (código {result.returncode}). Confira a instalação do OCR/PDF.')
    return result.stdout


def find_tesseract():
    candidates=[os.environ.get('CREDPLUS_TESSERACT',''),str(ASSETS/'ocr'/'tesseract.exe')]
    if getattr(sys,'frozen',False):candidates.append(str(Path(sys.executable).parent/'ocr'/'tesseract.exe'))
    candidates += [shutil.which('tesseract') or '',str(Path(os.environ.get('ProgramFiles','C:/Program Files'))/'Tesseract-OCR'/'tesseract.exe')]
    return next((p for p in candidates if p and Path(p).is_file()),None)


def _layout_words(page):
    # Pinned PyMuPDF layout renderer preserves table columns, unlike reading order.
    import io
    from pymupdf.__main__ import page_layout
    out=io.BytesIO()
    page_layout(page,out,2,3,False,False,
                flags=fitz.TEXT_PRESERVE_LIGATURES | fitz.TEXT_PRESERVE_WHITESPACE)
    return out.getvalue().decode('utf-8')


def read_pdf_text(path,progress=None,clean_ocr=False,ocr_psm=6,ocr_dpi=300):
    path=Path(path)
    texts=[];methods=[]
    with fitz.open(path) as doc:
        if doc.needs_pass: raise ValueError('PDF protegido por senha. Forneça uma cópia autorizada sem senha.')
        if len(doc)>60:raise ValueError('Limite de 60 páginas por proposta.')
        with tempfile.TemporaryDirectory(prefix='credplus-') as tmp:
            for i,page in enumerate(doc):
                if progress:progress(f'Página {i+1}/{len(doc)} — análise local…')
                text=_layout_words(page)
                if len(text.strip())<50:
                    exe=find_tesseract()
                    if not exe:raise RuntimeError('Este PDF é escaneado. O OCR Tesseract precisa ser incluído no pacote autorizado pela TI. Você também pode preencher os campos manualmente.')
                    img=Path(tmp)/f'page-{i}.png'
                    scale=min(ocr_dpi/72,5000/max(page.rect.width,page.rect.height))
                    page.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False).save(img)
                    if clean_ocr:
                        # Allianz colored table bands can become solid black under global OCR thresholding.
                        # Keep dark lettering and remove the light blue background on the temporary raster.
                        from PIL import Image
                        with Image.open(img) as source:
                            source.convert('L').point(lambda n:0 if n<140 else 255).save(img)
                    command=[exe,str(img),'stdout','-l','por+eng','--psm',str(ocr_psm),'-c','preserve_interword_spaces=1']
                    tessdata=os.environ.get('CREDPLUS_TESSDATA')
                    if not tessdata and (Path(exe).parent/'tessdata').is_dir():
                        tessdata=str(Path(exe).parent/'tessdata')
                    if tessdata:command += ['--tessdata-dir',tessdata]
                    text=_run(command).decode('utf-8');methods.append('OCR')
                else:methods.append('texto')
                texts.append(text)
    return '\n\f\n'.join(texts),methods


def extract_pdf(path, insurer, offer=1, progress=None):
    if insurer not in INSURERS or not 1<=offer<=len(INSURERS[insurer]): raise ValueError('Selecione a seguradora e a opção corretas.')
    rules=json.loads((ASSETS/'regras.json').read_text(encoding='utf-8'))
    full_text,methods=read_pdf_text(path,progress)
    results=extract(full_text,insurer,offer,rules,corrected=True)
    if insurer=='Allianz' and 'OCR' in methods:
        from .tables import allianz_ocr_table
        results.update(allianz_ocr_table(full_text,offer))
    # Retain only approved business fields; never expose complete source text in outputs.
    results={k:v for k,v in results.items() if k in FIELD_MAP}
    for key,item in results.items():
        try:item['display']=format_value(FIELD_MAP[key],item.get('value'))
        except ValueError:item.update(display='',state='Conferir formato')
        # All OCR and text candidates require human confirmation, including seemingly valid IDs.
        item['state']=('Conferir leitura - '+item.get('state','Revisar') if item.get('value') is not None else item.get('state','Não localizado'))
    return {'fields':results,'pages':len(methods),'methods':methods,'insurer':insurer,'offer':offer}


def export_xlsx(path,values,consultant='',phone=''):
    import openpyxl
    from openpyxl.styles import Alignment
    clean=validate_values(values)
    w=openpyxl.load_workbook(ASSETS/'modelo_limpo.xlsx');s=w.active
    overflow=[]
    for f in FIELDS:
        if f.cell:
            # Force string type to prevent formula injection from OCR/manual text.
            c=s[f.cell];value=clean[f.key] or 'Não informado'
            if len(value)>220:
                overflow.append((f.label,value));value='Detalhes na aba Complemento'
            c.value=value;c.data_type='s'
            c.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    s['B7']='Sicoob Credplus'
    s['A49']=f'Consultor / Gerente: {safe_text(consultant) or "__________________"} | WhatsApp: {safe_text(phone) or "__________________"}'
    s['A49'].data_type='s'
    # Model has no slots for these allowed fields: include an explicit continuation sheet.
    extra=w.create_sheet('Complemento')
    extra.append(['INFORMAÇÕES COMPLEMENTARES','Valor'])
    for f in FIELDS:
        if not f.cell:
            extra.append([f.label,clean[f.key] or 'Não informado'])
    for label,value in overflow:
        extra.append([label,value])
    extra.append(['Leitura','Campos não localizados não significam ausência de cobertura.'])
    for row in extra:
        for c in row:
            c.data_type='s';c.alignment=Alignment(wrap_text=True,vertical='top')
    extra.column_dimensions['A'].width=48;extra.column_dimensions['B'].width=90
    import textwrap
    for row in range(1,extra.max_row+1):
        value=str(extra.cell(row,2).value or '')
        lines=sum(max(1,len(textwrap.wrap(line,75))) for line in value.split('\n'))
        extra.row_dimensions[row].height=max(36,min(409,lines*15+8))
    extra.sheet_properties.pageSetUpPr.fitToPage=True;extra.page_setup.fitToWidth=1;extra.page_setup.fitToHeight=0
    s.row_dimensions[12].height=55
    for row in [10,11,15,27,28,29,31,39,40]:s.row_dimensions[row].height=48
    from openpyxl.utils import get_column_letter
    for row in s:
        needed=s.row_dimensions[row[0].row].height or 15
        for c in row:
            if c.value is None:continue
            width=s.column_dimensions[get_column_letter(c.column)].width or 13
            for merged in s.merged_cells.ranges:
                if merged.min_row==c.row and merged.min_col==c.column:
                    width=sum(s.column_dimensions[get_column_letter(col)].width or 13 for col in range(merged.min_col,merged.max_col+1));break
            lines=sum(max(1,len(textwrap.wrap(line,max(8,int(width*.82))))) for line in str(c.value).split('\n'))
            needed=max(needed,lines*14+8)
        s.row_dimensions[row[0].row].height=min(409,needed)
    s.page_setup.fitToHeight=0;s.print_title_rows='1:2'
    _atomic_save(path,lambda p:w.save(p))


def _atomic_save(path,writer):
    path=Path(path)
    fd,tmp=tempfile.mkstemp(prefix='.credplus-',suffix=path.suffix,dir=path.parent)
    os.close(fd)
    try:writer(tmp);os.replace(tmp,path)
    finally:Path(tmp).unlink(missing_ok=True)


def export_pdf(path,values,consultant='',phone=''):
    # Compatibility API for callers using a single set of fields.
    from .proposal import Proposal,Offer,COVER_KEYS,GLASS_KEYS
    from .document import export_document
    clean={key:format_value(FIELD_MAP[key],value) for key,value in values.items() if key in FIELD_MAP}
    commercial={k:v for k,v in clean.items() if k in (*COVER_KEYS,*GLASS_KEYS,'total') and v}
    common={k:v for k,v in clean.items() if k not in commercial and v}
    proposal=Proposal('',common,[Offer('Proposta',commercial)] if commercial else [])
    export_document(path,proposal,'pdf',consultant,phone)
